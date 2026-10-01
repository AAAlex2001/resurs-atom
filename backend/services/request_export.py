from datetime import timezone, timedelta
from io import BytesIO
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


def build_requests_docx(requests) -> bytes:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.2)
    section.left_margin = section.right_margin = Cm(1)
    document.styles['Normal'].font.name = 'Arial'
    document.styles['Normal'].font.size = Pt(9)
    document.styles['Normal'].paragraph_format.space_after = Pt(3)
    document.add_paragraph('Заявки с сайта', style='Title')
    document.add_paragraph(f'Всего заявок: {len(requests)}. Даты указаны по московскому времени.')
    table = document.add_table(rows=1, cols=9)
    table.style = 'Table Grid'
    table.autofit = False
    widths = [0.9, 2.4, 2.6, 3.5, 3.3, 3.4, 2.3, 6.5, 2.8]
    headers = ['№', 'Имя', 'Телефон', 'Email', 'Деятельность', 'Компания', 'ИНН', 'Сообщение', 'Дата']
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for request in requests:
        created = request.created_at
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        date = created.astimezone(timezone(timedelta(hours=3))).strftime('%d.%m.%Y, %H:%M')
        activity = getattr(request.activity, 'value', request.activity)
        values = [request.id, request.name, request.phone, request.email, activity,
                  request.company, request.inn, request.message, date]
        for cell, value, width in zip(table.add_row().cells, values, widths):
            cell.width = Cm(width)
            cell.text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value)) if value is not None else '—'
    output = BytesIO()
    document.save(output)
    return output.getvalue()
